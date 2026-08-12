from __future__ import annotations

import base64
import io
import json
import logging

import anthropic
from pydantic import ValidationError

from . import config
from .models import DriveDoc, Mietangebot

logger = logging.getLogger(__name__)

EXTRACTION_PROMPT = """Du liest ein Dokument aus dem Google Drive von Kromeich & Partner, einem Makler für Gewerbe- und Logistikimmobilien. Es soll ein MIETANGEBOT sein (oft "Indikatives Mietangebot"). Extrahiere die Konditionen für eine Vergleichsmieten-Datenbank.

Antworte ausschließlich mit einem JSON-Objekt nach dem Schema unten.

WAS IST EIN MIETANGEBOT: ein Dokument, das für eine konkrete Fläche konkrete Mietkonditionen nennt (Miete pro m² oder absolut, Laufzeit). Setze ist_mietangebot=false bei Exposés ohne Preis, Mietverträgen, Rechnungen, Flächengesuchen, Protokollen, reiner Korrespondenz.

SONDERFÄLLE – markiere sie statt sie zu erfinden:
- ist_vorlage=true: leere Vorlage/Muster/Blanko ohne echte Objektdaten.
- ist_eigenmiete=true: Rechnung/Angebot für die EIGENE Büromiete von Kromeich (Mieter ist die Kromeich GmbH selbst). Das ist kein Marktangebot.
- ist_anlage=true: Anlagen-/Beiblatt-Dokument (Grundrisse, AGB, Flächenaufstellung) ohne eigene Konditionen.

FELDER (null wenn nicht im Dokument – NICHTS schätzen oder ableiten):
- objekt: Objekt-/Projektname, z.B. "Sun Park Bitterfeld", "AUK Logis"
- adresse: Straße + Hausnummer
- plz: 5-stellige PLZ des OBJEKTS (nicht des Absenders!)
- ort: Ort des Objekts
- anbieter: wer die Fläche anbietet (Vermieter/Eigentümer bzw. dessen Makler), z.B. "Mileway", "HIH", "Westcore"
- empfaenger: an wen das Angebot gerichtet ist (Mietinteressent)
- datum: Datum des Angebots als "JJJJ-MM-TT". Steht kein Datum im Text, nimm null (nicht das heutige Datum).
- flaeche_qm: angebotene Gesamtfläche in m² als Zahl (Hallenfläche + Büro, wenn getrennt genannt: Summe)
- nutzungsart: kurz, z.B. "Logistik", "Halle", "Produktion", "Büro", "Freifläche"
- sicherheit: Mietsicherheit, z.B. "3 Monatsmieten Bürgschaft"
- indexierung: Wertsicherung, z.B. "VPI 100%, jährlich"
- eigenes_angebot: true, wenn Kromeich & Partner das Angebot ERSTELLT/VERSANDT hat; false, wenn K&P es von einem Dritten ERHALTEN hat; null wenn unklar.

OPTIONEN – das Herzstück. Gib EINE Option pro LAUFZEIT-Variante zurück:
- Nennt das Dokument je Laufzeit einen anderen Mietpreis (typische Laufzeitstaffel,
  z.B. "5 Jahre 4,50 €/m², 7 Jahre 4,30 €/m², 10 Jahre 4,05 €/m²"), dann ist das
  EINE Option JE Laufzeit – also drei Optionen mit je eigener laufzeit_monate.
- Steigt die Miete INNERHALB einer Laufzeit im Zeitverlauf (Staffel über die Jahre),
  dann ist das EINE Option: kaltmiete = die ANFANGSMIETE, und beschreibe den
  Verlauf in "hinweis".
- Nur ein Preis und eine Laufzeit: genau eine Option.
- Miete nur absolut genannt (z.B. "12.500 € monatlich")? Dann kaltmiete_absolut_eur
  füllen und kaltmiete_eur_qm auf null lassen – NICHT selbst umrechnen.
- Jahresmiete erkennbar (z.B. "150.000 € p.a.")? Auf den MONAT umrechnen und in
  kaltmiete_absolut_eur eintragen, den Ursprung in "hinweis" notieren.

Pro Option:
- laufzeit_monate: Laufzeit in MONATEN als ganze Zahl (5 Jahre -> 60)
- kaltmiete_eur_qm: Nettokaltmiete/Grundmiete in €/m²/MONAT als Zahl (Dezimalpunkt).
  Ist sie im Dokument pro Jahr angegeben, durch 12 teilen und in "hinweis" notieren.
- kaltmiete_absolut_eur: absolute MONATS-Nettokaltmiete, falls so genannt
- nebenkosten_eur_qm: Nebenkosten/Betriebskosten in €/m²/Monat
- mietfreie_monate: mietfreie Zeit in Monaten als Zahl
- flaeche_qm: nur wenn diese Option eine ANDERE Fläche betrifft als das Gesamtangebot
- hinweis: kurzer Zusatz (Staffelverlauf, Umrechnung, Teilfläche) – sonst null

Zahlen im deutschen Format umsetzen: "4,58" -> 4.58, "12.500" -> 12500, "40.000 qm" -> 40000.

Schema:
{{
  "ist_mietangebot": true,
  "objekt": "..." | null,
  "adresse": "..." | null,
  "plz": "..." | null,
  "ort": "..." | null,
  "anbieter": "..." | null,
  "empfaenger": "..." | null,
  "datum": "JJJJ-MM-TT" | null,
  "flaeche_qm": 0.0 | null,
  "nutzungsart": "..." | null,
  "sicherheit": "..." | null,
  "indexierung": "..." | null,
  "eigenes_angebot": true | false | null,
  "ist_vorlage": false,
  "ist_eigenmiete": false,
  "ist_anlage": false,
  "optionen": [
    {{
      "laufzeit_monate": 60 | null,
      "kaltmiete_eur_qm": 4.58 | null,
      "kaltmiete_absolut_eur": null,
      "nebenkosten_eur_qm": 2.15 | null,
      "mietfreie_monate": 3 | null,
      "flaeche_qm": null,
      "hinweis": null
    }}
  ],
  "confidence": 0.0,
  "begruendung": "1 kurzer Satz"
}}

Dateiname: {dateiname}
{ordner_hinweis}
"""


def _get_client() -> anthropic.Anthropic:
    return anthropic.Anthropic(api_key=config.anthropic_api_key())


def _docx_text(data: bytes) -> str | None:
    try:
        import docx  # type: ignore
    except ImportError:
        logger.warning("python-docx fehlt – DOCX wird übersprungen")
        return None
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _pptx_text(data: bytes) -> str | None:
    try:
        from pptx import Presentation  # type: ignore
    except ImportError:
        logger.warning("python-pptx fehlt – PPTX wird übersprungen")
        return None
    presentation = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for index, slide in enumerate(presentation.slides, 1):
        parts.append(f"--- Folie {index} ---")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
    return "\n".join(parts)


def _xlsx_text(data: bytes) -> str | None:
    try:
        import openpyxl  # type: ignore
    except ImportError:
        logger.warning("openpyxl fehlt – XLSX wird übersprungen")
        return None
    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"--- Blatt {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def build_content(doc: DriveDoc, data: bytes | None, text: str | None) -> list[dict] | None:
    """Baut die Anthropic-Content-Blöcke für ein Dokument.

    PDFs gehen als Dokument-Block direkt an Claude (liest auch gescannte
    Angebote und Tabellen-Layouts, die eine Textextraktion zerreißt).
    """
    ordner_hinweis = f"Drive-Ordner: {doc.pfad_hinweis}" if doc.pfad_hinweis else ""
    prompt = EXTRACTION_PROMPT.format(dateiname=doc.name, ordner_hinweis=ordner_hinweis)

    if doc.mime_type == config.MIME_PDF and data:
        if len(data) > config.MAX_PDF_BYTES:
            logger.warning("%s ist %d bytes – über dem PDF-Limit, übersprungen", doc.name, len(data))
            return None
        return [
            {
                "type": "document",
                "source": {
                    "type": "base64",
                    "media_type": "application/pdf",
                    "data": base64.standard_b64encode(data).decode("ascii"),
                },
            },
            {"type": "text", "text": prompt},
        ]

    if text is None and data is not None:
        if doc.mime_type == config.MIME_DOCX:
            text = _docx_text(data)
        elif doc.mime_type == config.MIME_PPTX:
            text = _pptx_text(data)
        elif doc.mime_type == config.MIME_XLSX:
            text = _xlsx_text(data)

    if not text or not text.strip():
        return None

    return [{"type": "text", "text": f"{prompt}\n\nDokumentinhalt:\n{text[:config.MAX_TEXT_CHARS]}"}]


def parse_extraction_json(text: str) -> dict:
    """```-Fences strippen, erste {...} bis letzte } extrahieren, json.loads."""
    text = text.strip()
    if text.startswith("```"):
        text = text.split("\n", 1)[1].rsplit("```", 1)[0].strip()
    start = text.find("{")
    end = text.rfind("}")
    if start != -1 and end != -1 and end > start:
        text = text[start : end + 1]
    return json.loads(text)


def extract_angebot(content: list[dict]) -> Mietangebot | None:
    """Schickt die Content-Blöcke an Claude. None bei API-/Parse-Fehler."""
    try:
        response = _get_client().messages.create(
            model=config.CLAUDE_MODEL,
            max_tokens=4096,
            messages=[{"role": "user", "content": content}],
        )
        raw = parse_extraction_json(response.content[0].text)
        return Mietangebot.model_validate(raw)
    except anthropic.APIError as e:
        logger.error("Claude API error bei Angebots-Extraktion: %s", e)
        return None
    except (json.JSONDecodeError, ValidationError, KeyError, IndexError) as e:
        logger.error("Extraktions-Antwort nicht parsebar: %s", e)
        return None
